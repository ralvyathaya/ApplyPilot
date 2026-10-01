import csv
import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from applypilot.export import (
    _calculate_freshness,
    _determine_status,
    _format_date,
    _parse_date_safe,
    export_jobs,
    export_to_csv,
    export_to_xlsx,
    fetch_jobs_for_export,
)


def test_parse_date_safe():
    assert _parse_date_safe(None) is None
    assert _parse_date_safe("") is None

    # ISO format
    dt1 = _parse_date_safe("2026-09-28T07:08:55+00:00")
    assert dt1 is not None
    assert dt1.year == 2026 and dt1.month == 9 and dt1.day == 28

    # YYYY-MM-DD
    dt2 = _parse_date_safe("2026-09-25")
    assert dt2 is not None
    assert dt2.year == 2026 and dt2.month == 9 and dt2.day == 25

    # RFC 2822
    dt3 = _parse_date_safe("Mon, 28 Sep 2026 11:01:08 +0000")
    assert dt3 is not None
    assert dt3.year == 2026 and dt3.month == 9 and dt3.day == 28


def test_calculate_freshness():
    ref_time = datetime.datetime(2026, 9, 29, 12, 0, 0, tzinfo=datetime.timezone.utc)

    # Missing date
    assert _calculate_freshness(None, ref_time=ref_time) == "-"

    # < 24h (e.g. 5 hours ago)
    t_24h = "2026-09-29T07:00:00+00:00"
    assert _calculate_freshness(t_24h, ref_time=ref_time) == "< 24 Hours"

    # 1 - 3 days (e.g. 2 days ago)
    t_2d = "2026-09-27T12:00:00+00:00"
    assert _calculate_freshness(t_2d, ref_time=ref_time) == "1 - 3 Days"

    # Within 1 week (e.g. 5 days ago)
    t_5d = "2026-09-24T12:00:00+00:00"
    assert _calculate_freshness(t_5d, ref_time=ref_time) == "Within 1 Week"

    # > 1 month (e.g. 40 days ago)
    t_40d = "2026-08-20T12:00:00+00:00"
    assert _calculate_freshness(t_40d, ref_time=ref_time) == "> 1 Month"


def test_determine_status():
    assert _determine_status({"applied_at": "2026-09-20", "apply_status": "applied"}) == "Applied"
    assert _determine_status({"applied_at": "2026-09-20", "apply_status": "failed"}) == "Apply Failed"
    assert _determine_status({"apply_status": "expired"}) == "Expired"
    assert _determine_status({
        "tailored_resume_path": "/path/resume.txt",
        "cover_letter_path": "/path/cl.txt",
    }) == "Ready to Apply"
    assert _determine_status({"tailored_resume_path": "/path/resume.txt"}) == "Tailored"
    assert _determine_status({"fit_score": 8}) == "Scored"
    assert _determine_status({"full_description": "Job desc"}) == "Enriched"
    assert _determine_status({}) == "Discovered"


def test_export_to_csv(tmp_path):
    mock_jobs = [
        {
            "url": "https://example.com/job1",
            "title": "Software Engineer",
            "company": "Acme Corp",
            "location": "Jakarta",
            "site": "linkedin",
            "salary": "$2000",
            "posted_at": "2026-09-28",
            "discovered_at": "2026-09-28T10:00:00",
            "fit_score": 9,
            "score_reasoning": "Strong match",
            "tailored_resume_path": "/path/to/resume1.txt",
            "cover_letter_path": "/path/to/cl1.txt",
            "applied_at": None,
            "apply_status": None,
        }
    ]

    out_csv = tmp_path / "test_export.csv"
    export_to_csv(mock_jobs, out_csv)

    assert out_csv.exists()
    with open(out_csv, "r", encoding="utf-8-sig") as f:
        reader = list(csv.reader(f))
        assert len(reader) == 2
        headers = reader[0]
        assert "Fit Score" in headers
        assert "Job Title" in headers
        assert "Freshness" in headers

        first_row = reader[1]
        assert first_row[0] == "9"
        assert first_row[3] == "Software Engineer"
        assert first_row[4] == "Acme Corp"


def test_export_to_xlsx(tmp_path):
    mock_jobs = [
        {
            "url": "https://example.com/job1",
            "title": "Lead Developer",
            "company": "Tech Inc",
            "location": "Remote",
            "site": "glints",
            "salary": None,
            "posted_at": "2026-09-29T10:00:00",
            "discovered_at": "2026-09-29T11:00:00",
            "fit_score": 8,
            "score_reasoning": "Excellent experience",
            "tailored_resume_path": "resume.pdf",
            "cover_letter_path": None,
            "applied_at": "2026-09-29T12:00:00",
            "apply_status": "applied",
        }
    ]

    out_xlsx = tmp_path / "test_export.xlsx"
    with patch("applypilot.export.get_stats", return_value={"total": 1, "scored": 1, "applied": 1}):
        export_to_xlsx(mock_jobs, out_xlsx)

    assert out_xlsx.exists()
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Jobs" in wb.sheetnames
    assert "Summary Stats" in wb.sheetnames

    ws_jobs = wb["Jobs"]
    assert ws_jobs.cell(row=1, column=1).value == "Fit Score"
    assert ws_jobs.cell(row=2, column=1).value == 8
    assert ws_jobs.cell(row=2, column=4).value == "Lead Developer"
    assert ws_jobs.cell(row=2, column=8).value == "Open Link"
    assert ws_jobs.cell(row=2, column=8).hyperlink.target == "https://example.com/job1"


def test_export_jobs_invalid_format():
    with pytest.raises(ValueError, match="Unsupported format"):
        export_jobs(fmt="docx", auto_open=False)
